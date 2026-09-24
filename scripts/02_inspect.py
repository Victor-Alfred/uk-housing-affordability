"""
Inspect an ONS workbook's structure without opening it in a spreadsheet app.

This is the discovery spike in script form. It answers the questions you
would otherwise answer by eye - how many sheets, where the headers start,
how years are labelled, what suppressed cells contain - and it answers
them reproducibly, so the output can be pasted straight into
docs/source-assessment.md.

Usage:
    python scripts/02_inspect.py <file>
    python scripts/02_inspect.py <file> --sheet 6c --headers
    python scripts/02_inspect.py <file> --sheet 6c --skip-cols 4 --marker-map
    python scripts/02_inspect.py <file> --sheet 6c --skip-cols 4 --find "Bournemouth"
    python scripts/02_inspect.py <file> --sheet 6c --skip-cols 4 --markers --rows 3
"""

import argparse
import sys
from pathlib import Path

from openpyxl import load_workbook

WIDTH = 22          # max characters shown per cell
MAX_COLS = 10       # columns shown per row


def cell_text(value):
    if value is None:
        return ""
    text = str(value).replace("\n", " ").strip()
    return text[:WIDTH - 1] + "…" if len(text) > WIDTH else text


def show_headers(ws, header_row):
    """Print the full header row with column indices.

    ONS sheets vary in column count even within a table group, so the
    parser must read year columns from the header rather than assume a
    fixed width. This is how you find out what to read.
    """
    print(f"\n{'=' * 78}")
    print(f"SHEET: {ws.title}   header row {header_row}   "
          f"({ws.max_row} rows x {ws.max_column} cols)")
    print("=" * 78)
    for row in ws.iter_rows(min_row=header_row, max_row=header_row,
                            values_only=True):
        for i, value in enumerate(row, start=1):
            print(f"  col {i:>3}: "
                  f"{cell_text(value) if value is not None else '(empty)'}")


def find_rows(ws, needle, header_row, label_cols):
    """Print full rows whose label columns contain `needle`."""
    print(f"\n{'=' * 78}")
    print(f"SHEET: {ws.title}   rows matching {needle!r}")
    print("=" * 78)
    header = next(ws.iter_rows(min_row=header_row, max_row=header_row,
                               values_only=True))
    hits = 0
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        labels = " ".join(str(v) for v in row[:label_cols] if v is not None)
        if needle.lower() not in labels.lower():
            continue
        hits += 1
        print(f"\n  {labels}")
        for h, v in zip(header[label_cols:], row[label_cols:]):
            if h is not None:
                print(f"    {cell_text(h):<24} {cell_text(v)}")
    if not hits:
        print("  no matching rows")


def markers_by_column(ws, header_row, label_cols):
    """Count non-numeric body values per column.

    A total count says how much is missing; this says WHERE. Gaps confined
    to the earliest years are a historical artefact. Gaps inside the
    analysis window are a threat to the analysis.
    """
    header = next(ws.iter_rows(min_row=header_row, max_row=header_row,
                               values_only=True))
    counts = {}
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        for h, v in zip(header[label_cols:], row[label_cols:]):
            if h is None or v is None or isinstance(v, (int, float)):
                continue
            counts[str(h)] = counts.get(str(h), 0) + 1

    print(f"\n{'=' * 78}")
    print(f"SHEET: {ws.title}   non-numeric values by column")
    print("=" * 78)
    total = 0
    for h in header[label_cols:]:
        if h is None:
            continue
        n = counts.get(str(h), 0)
        total += n
        print(f"  {str(h):<24} {n:>4}  {'#' * min(n, 50)}")
    print(f"  {'TOTAL':<24} {total:>4}")


def show_sheet(ws, rows):
    print(f"\n{'=' * 78}")
    print(f"SHEET: {ws.title}   ({ws.max_row} rows x {ws.max_column} cols)")
    print("=" * 78)

    for i, row in enumerate(ws.iter_rows(max_row=rows, max_col=MAX_COLS,
                                         values_only=True), start=1):
        cells = " | ".join(f"{cell_text(v):<{WIDTH}}" for v in row)
        print(f"{i:>3} | {cells}")

    if ws.max_column > MAX_COLS:
        print(f"    ... {ws.max_column - MAX_COLS} further column(s) not shown")


def scan_markers(ws, skip_rows, skip_cols):
    """Non-numeric values in the body of a sheet - i.e. suppression markers.

    ONS suppresses small-count cells. Knowing the exact marker matters:
    how you treat it (null, zero, excluded) is a business rule that
    belongs in the decision log, not buried in a parser.
    """
    found = {}
    for row in ws.iter_rows(min_row=skip_rows + 1, values_only=True):
        for value in row[skip_cols:]:      # skip geography code/name columns
            if value is None or isinstance(value, (int, float)):
                continue
            text = str(value).strip()
            if not text or len(text) > 12:   # long strings are labels
                continue
            try:
                float(text.replace(",", ""))
            except ValueError:
                found[text] = found.get(text, 0) + 1
    return found


def main():
    p = argparse.ArgumentParser()
    p.add_argument("path", type=Path)
    p.add_argument("--rows", type=int, default=10,
                   help="rows to show per sheet (default 10)")
    p.add_argument("--sheet", help="only sheets whose name contains this")
    p.add_argument("--markers", action="store_true",
                   help="report non-numeric values in sheet bodies")
    p.add_argument("--skip-rows", type=int, default=2,
                   help="header rows to skip when scanning markers")
    p.add_argument("--skip-cols", type=int, default=2,
                   help="leading label columns to skip when scanning markers")
    p.add_argument("--headers", action="store_true",
                   help="print the full header row with column indices")
    p.add_argument("--header-row", type=int, default=2,
                   help="which row holds the column headers (default 2)")
    p.add_argument("--find", help="print full rows whose labels contain this")
    p.add_argument("--marker-map", action="store_true",
                   help="count non-numeric body values per column")
    args = p.parse_args()

    if not args.path.exists():
        sys.exit(f"not found: {args.path}")

    # read_only keeps memory flat on large workbooks; data_only returns
    # cached formula results rather than the formulas themselves.
    wb = load_workbook(args.path, read_only=True, data_only=True)

    print(f"FILE: {args.path.name}")
    print(f"{len(wb.sheetnames)} sheet(s):\n")
    for name in wb.sheetnames:
        print(f"  - {name}")

    targets = [n for n in wb.sheetnames
               if not args.sheet or args.sheet.lower() in n.lower()]
    if args.sheet and not targets:
        sys.exit(f"\nno sheet matching {args.sheet!r}")

    for name in targets:
        ws = wb[name]
        if args.headers:
            show_headers(ws, args.header_row)
            continue
        if args.find:
            find_rows(ws, args.find, args.header_row, args.skip_cols)
            continue
        if args.marker_map:
            markers_by_column(ws, args.header_row, args.skip_cols)
            continue
        show_sheet(ws, args.rows)
        if args.markers:
            markers = scan_markers(ws, args.skip_rows, args.skip_cols)
            if markers:
                print("\n  non-numeric values in body:")
                for text, count in sorted(markers.items(),
                                          key=lambda kv: -kv[1])[:10]:
                    print(f"    {text!r:<16} x{count}")

    wb.close()


if __name__ == "__main__":
    main()